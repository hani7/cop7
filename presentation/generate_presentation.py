"""Build the French presentation as editable PowerPoint and matching vector PDF.

Uses isolated python-pptx in .runtime/presentation-tools and system ReportLab.
No application data, credentials, or application source is modified.
"""
from pathlib import Path
import sys
import json

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / '.runtime' / 'presentation-tools'))
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from PIL import Image

W, H = 960, 540
PETROL = '083F48'
GREEN = '087F73'
MINT = 'B8EAD9'
GOLD = 'D6B76D'
BG = 'F2F6F5'
WHITE = 'FFFFFF'
INK = '173E46'
MUTED = '557077'
LINE = 'D9E5E2'
LOGO = ROOT / 'frontend/src/assets/CO-1-1-1.png'

SLIDES = []

def add(section, title, subtitle, cards=None, takeaway='', notes='', sources='', kind='cards', **kw):
    SLIDES.append(dict(section=section, title=title, subtitle=subtitle, cards=cards or [],
                       takeaway=takeaway, notes=notes, sources=sources, kind=kind, **kw))

add('PRÉSENTATION DU PROJET', 'COPETROLITE',
    'La gestion des interventions industrielles', kind='cover',
    takeaway='Les équipes. Le terrain. La validation. Le suivi financier.',
    notes='Présenter Copetrolite comme une application de gestion des interventions industrielles. Elle relie les entreprises clientes, les équipes internes et les techniciens dans un circuit commun. Cette présentation décrit le code présent dans le projet au 18 octobre 2026 ; elle ne constitue pas une preuve de déploiement en production.')

add('01 / VISION', 'Un suivi continu de chaque intervention',
    'Centraliser les informations et clarifier les responsabilités.', [
    ('ORGANISER', 'Réunir les clients, les projets, les techniciens et leurs documents dans une même application.'),
    ('SUIVRE LE TERRAIN', 'Permettre au technicien de déclarer ses heures et de décrire les travaux réalisés.'),
    ('FAIRE VALIDER', 'Confier au responsable client autorisé la validation ou le refus motivé des pointages.'),
    ('PRÉPARER LA FACTURATION', 'Transformer les heures validées en factures, puis suivre les règlements et les soldes.')],
    'Une même intervention reste identifiable du premier pointage au règlement.',
    'Les bénéfices présentés sont les objectifs du fonctionnement de l’application. Aucun gain de productivité chiffré n’a été mesuré. Le socle existant couvre le suivi opérationnel, documentaire et financier.', 'Modèles, services métier et interface')

add('01 / VISION', 'Le parcours métier de bout en bout',
    'Chaque étape possède un responsable et une action définie.', kind='flow',
    steps=[('PRÉPARER', 'Admin / superviseur', 'Client, projet et affectations'),
           ('DÉCLARER', 'Technicien', 'Date, durée et travaux'),
           ('VALIDER', 'Responsable client', 'Accord ou refus motivé'),
           ('FACTURER', 'Administrateur', 'Pointages validés'),
           ('ENCAISSER', 'Administrateur', 'Paiement et solde')],
    cards=[('UNE VALIDATION CLIENT', 'L’administrateur et le superviseur suivent les heures ; la décision de validation appartient au responsable client autorisé.'),
           ('UN HISTORIQUE CONSERVÉ', 'Les changements d’état sont datés et associés à leur auteur. Les corrections et les refus restent traçables.')],
    takeaway='La responsabilité de chaque intervenant est visible dans le circuit.',
    notes='Le pointage est d’abord un brouillon. Sa soumission déclenche le travail de validation client. Seul un pointage validé et encore non facturé peut entrer dans une facture. Le paiement constitue une étape distincte de la facturation.', sources='eliteapp/access.py · services.py · serializers.py')

add('01 / VISION', 'Les modules disponibles',
    'Une interface en français, adaptée au rôle connecté.', [
    ('PILOTAGE', 'Vue d’ensemble, projets en cours, heures pointées, pointages à valider et recherche dans les listes.'),
    ('ORGANISATION', 'Clients, projets, techniciens, comptes de l’équipe et affectations aux interventions.'),
    ('DOSSIERS TECHNICIENS', 'Identité, spécialité, expérience, photo, CV et documents de certification.'),
    ('SUIVI FINANCIER', 'Tarifs par projet, calendrier des jours fériés, factures et paiements clients ; gains personnels côté technicien.')],
    'Les menus et les données visibles dépendent des autorisations.',
    'Pour l’administrateur, Équipe et Calendrier se trouvent dans Paramètres. Les autres menus principaux sont Vue d’ensemble, Projets, Pointages, Techniciens, Clients, Facturation et Paiements. Le contenu de Paiements diffère entre administrateur et technicien.', 'frontend/src/App.jsx · forms.jsx')

add('02 / LES ACCÈS', 'Quatre rôles, des responsabilités distinctes',
    '« Équipe » est un menu de gestion des personnes ; le rôle interne est « Superviseur interne ».', [
    ('ADMINISTRATEUR', 'Gère les comptes, les dossiers, les accès client, les clients, les projets, les factures et les paiements.'),
    ('SUPERVISEUR INTERNE', 'Organise les clients et les projets, affecte les techniciens et consulte les pointages et dossiers internes.'),
    ('TECHNICIEN', 'Consulte ses projets, déclare ses propres heures, corrige les pointages autorisés et consulte ses gains.'),
    ('CHEF DE PROJET CLIENT', 'Consulte son périmètre de projets et valide ou refuse les pointages soumis. Un compte dédié peut être limité à un projet.')],
    'L’accès est contrôlé par le serveur, au-delà des menus affichés.',
    'Le code définit exactement quatre rôles. Il n’existe pas de rôle comptable séparé. Un compte d’accès client par projet utilise le rôle de chef de projet client avec un périmètre plus restreint.', 'eliteapp/models.py · access.py · views.py')

add('02 / LES ACCÈS', 'Matrice des droits',
    'Synthèse des opérations et des périmètres autorisés.', kind='matrix',
    headers=['Fonction', 'Admin', 'Superviseur', 'Technicien', 'Client'],
    rows=[['Comptes de l’équipe', 'Gérer', 'Consulter¹', '—', 'Consulter¹'],
          ['Accès client par projet', 'Gérer', '—', '—', '—'],
          ['Clients / projets', 'Gérer', 'Gérer', 'Ses projets', 'Ses projets'],
          ['CV / photos / certificats', 'Gérer', 'Consulter', '—', '—'],
          ['Pointages', 'Consulter', 'Consulter', 'Saisir / soumettre²', 'Consulter³'],
          ['Validation / refus', '—', '—', '—', 'Son périmètre'],
          ['Tarifs commerciaux', 'Gérer', 'Gérer', '—', '—'],
          ['Factures / encaissements', 'Gérer', '—', '—', '—'],
          ['Gains technicien', 'Coûts projet', 'Coûts projet', 'Ses gains', '—'],
          ['Calendrier férié', 'Gérer', '—', '—', '—']],
    takeaway='¹ Selon périmètre.  ² Ses propres pointages.  ³ Les brouillons sont exclus.',
    notes='Le client peut consulter des informations limitées sur les personnes de son périmètre, sans documents privés. Le superviseur consulte les comptes mais ne les crée ni ne les modifie. Le technicien retrouve ses gains dans un endpoint distinct, sans accès aux factures ou paiements clients. Les coûts de projet sont consultables en interne.',
    sources='Contrôles serveur : access.py · views.py · serializers.py')

add('03 / ADMINISTRATION', 'Le tableau de bord administrateur',
    'Identifier rapidement l’activité et les actions en attente.', [
    ('ACTIVITÉ', 'Heures pointées et projets en cours pour suivre le volume d’interventions enregistré.'),
    ('VALIDATIONS', 'Nombre de pointages soumis en attente de décision du responsable client.'),
    ('TRÉSORERIE CLIENT', 'Reste à encaisser calculé à partir du solde des factures émises.'),
    ('ACCÈS AUX DÉTAILS', 'Projets actifs et derniers pointages ; ouverture des listes pour consulter les informations complètes.')],
    'Les indicateurs des autres rôles sont limités à leurs données accessibles.',
    'Le tableau de bord actuel présente quatre indicateurs principaux. Pour les profils non administrateurs, les heures validées remplacent le reste à encaisser. L’indicateur heures validées exclut les pointages déjà passés à l’état facturé ; il ne doit donc pas être présenté comme un cumul historique de toute l’activité approuvée.', 'DashboardView · frontend/src/App.jsx')

add('03 / ADMINISTRATION', 'Clients : un référentiel structuré',
    'Administrateur et superviseur renseignent les entreprises clientes.', [
    ('IDENTITÉ ET CONTACT', 'Société, personne à contacter, téléphone, email et logo de l’entreprise.'),
    ('INFORMATIONS ADMINISTRATIVES', 'N° RC, NIF, NIS, n° article, adresse de facturation et date de début de travail.'),
    ('LIEN AVEC LES PROJETS', 'Chaque projet est rattaché à un client. Les comptes responsables client sont rattachés à leur entreprise.'),
    ('CONSERVATION DE L’HISTORIQUE', 'Désactivation possible. La suppression, réservée à l’administrateur, est bloquée si des projets ou comptes sont liés.')],
    'Une fiche client sert de point de référence aux projets et à la facturation.',
    'La fiche client regroupe les coordonnées nécessaires au suivi. Le formulaire actuel rend plusieurs champs administratifs obligatoires. Une facture conserve un instantané de certaines informations client à son émission.', 'ClientViewSet · ClientSerializer · forms.jsx')

add('03 / ADMINISTRATION', 'Équipe : créer et maîtriser les comptes',
    'Parcours administrateur : Paramètres → Équipe.', [
    ('CRÉER UN COLLABORATEUR', 'Renseigner l’identité, l’identifiant, le mot de passe, le rôle et le statut actif.'),
    ('RATTACHER LE CLIENT', 'Un chef de projet client doit être lié à une entreprise cliente. Les autres rôles n’ont pas ce rattachement.'),
    ('GÉRER LE CYCLE DE VIE', 'Modifier les informations, remplacer un mot de passe et désactiver un compte lorsque l’accès doit cesser.'),
    ('PRÉSERVER LES DROITS', 'Le superviseur et le client consultent les personnes autorisées. La création et la modification des comptes restent administratives.')],
    'Les techniciens disposent aussi d’un menu dédié pour leur dossier complet.',
    'Le mot de passe est enregistré sous forme hachée. Les écrans ne permettent pas de relire le mot de passe stocké. Les restrictions de modification empêchent notamment certaines incohérences de rattachement lorsqu’un historique existe.', 'UserViewSet · UserSerializer · forms.jsx')

add('03 / ADMINISTRATION', 'Le dossier complet du technicien',
    'Centraliser les compétences, les pièces et les affectations.', [
    ('IDENTITÉ', 'Nom, prénom, date de naissance, adresse, téléphone, email, photo et date d’affectation.'),
    ('COMPÉTENCES', 'Poste ou spécialité, expérience professionnelle avec mise en forme et informations de certification.'),
    ('DOCUMENTS', 'CV et plusieurs certificats en pièces jointes, avec consultation et téléchargement autorisés.'),
    ('GESTION', 'Matricule généré automatiquement, recherche, filtre actif / désactivé et consultation des projets affectés.')],
    'L’administrateur modifie le dossier ; le superviseur peut consulter les documents.',
    'Le matricule est généré selon un compteur annuel au format CP-année-numéro. La photo, l’adresse, le CV et les autres informations privées ne sont pas exposés aux clients ou aux techniciens par la réponse utilisateur limitée. Une suppression de technicien avec projets ou pointages est bloquée.', 'UserSerializer · UserViewSet · TechnicianPhoto.jsx · CertificateUploads.jsx')

add('04 / CV & PHOTO', 'Détecter l’expérience depuis un CV',
    'Une aide à la saisie intégrée à la fiche technicien.', [
    ('IMPORT DU DOCUMENT', 'Sélectionner un CV PDF, JPG ou PNG, de 10 Mo maximum. Un CV enregistré peut aussi être relu.'),
    ('LECTURE DU TEXTE', 'Lecture directe des PDF contenant du texte ; recours à l’OCR local Windows pour les images et pages scannées.'),
    ('REPÉRAGE DE L’EXPÉRIENCE', 'Recherche de rubriques en français ou en anglais et prise en compte de certaines mises en page à colonnes.'),
    ('RELECTURE HUMAINE', 'L’expérience proposée est modifiable. Si la rubrique n’est pas reconnue, le texte extrait permet de choisir les passages utiles.')],
    'Le texte est proposé pour vérification avant enregistrement.',
    'L’extraction repose sur PyMuPDF et des règles de repérage des rubriques, avec OCR Windows si nécessaire. Les titres reconnus comprennent expérience professionnelle, parcours professionnel et work experience. Il ne s’agit pas d’une analyse universelle de toutes les mises en page. Le résultat retourné est limité à 20 000 caractères.', 'eliteapp/cv.py · CVImport.jsx')

add('04 / CV & PHOTO', 'De l’import à une expérience exploitable',
    'Parcours administrateur : Techniciens → Ajouter ou Modifier.', kind='flow',
    steps=[('CHOISIR', 'CV local', 'PDF, JPG ou PNG'),
           ('EXTRAIRE', 'Bouton dédié', 'Lire l’expérience'),
           ('RELIRE', 'Administrateur', 'Vérifier et corriger'),
           ('METTRE EN FORME', 'Éditeur intégré', 'Paragraphes et listes'),
           ('ENREGISTRER', 'Fiche technicien', 'Sauvegarder le dossier')],
    cards=[('SI LE CHAMP EST VIDE', 'Une expérience reconnue remplit le champ. L’administrateur relit le texte avant de sauvegarder.'),
           ('SI UNE EXPÉRIENCE EXISTE', 'Un aperçu est proposé. Le remplacement du contenu se fait par une action explicite de l’administrateur.')],
    takeaway='« Extraire » prépare le contenu ; « Enregistrer » sauvegarde le CV et la fiche.',
    notes='La distinction entre extraction et sauvegarde est essentielle lors de la démonstration. Le bouton d’extraction ne suffit pas pour joindre le CV. L’éditeur propose gras, italique, soulignement, listes et annuler/rétablir. En cas d’absence de rubrique détectée, le texte proposé peut être corrigé puis adopté.', sources='CVImport.jsx · ExperienceEditor.jsx · UserSerializer')

add('04 / CV & PHOTO', 'Extraire la photo du CV',
    'Un aperçu visuel pour enrichir rapidement le profil.', [
    ('ACTION UTILISATEUR', 'Cliquer sur « Utiliser la photo du CV » depuis un CV sélectionné ou déjà enregistré.'),
    ('SÉLECTION DE L’IMAGE', 'Le système recherche une image candidate selon sa taille, ses proportions et sa place dans le document.'),
    ('APERÇU ET SAUVEGARDE', 'L’image proposée apparaît dans la fiche. Vérifier qu’il s’agit bien du portrait, puis cliquer sur Enregistrer.'),
    ('PHOTO MANUELLE PRIORITAIRE', 'Une photo peut être importée séparément : PNG, WebP, JPG/JPEG ou AVIF, de 5 Mo maximum.')],
    'La sélection est heuristique : un logo ou une page scannée peut nécessiter un remplacement.',
    'Le code sélectionne une image candidate ; il ne réalise pas de reconnaissance faciale et ne garantit pas un recadrage du visage sur un CV scanné. Lors de la sauvegarde d’une fiche sans photo, une image peut aussi être extraite automatiquement du CV. Une photo déjà présente est conservée ; une photo explicitement importée est prioritaire.', 'extract_cv_photo · save_photo_from_cv · CVImport.jsx · TechnicianPhoto.jsx')

add('04 / CV & PHOTO', 'CV : qualité, confidentialité et contrôle',
    'Le dossier reste vérifiable par l’équipe autorisée.', [
    ('QUALITÉ DU RÉSULTAT', 'La netteté du scan, les colonnes et les titres du CV influencent l’extraction. La saisie manuelle reste disponible.'),
    ('TRAITEMENT LOCAL', 'Le texte et la photo sont traités localement. Le parcours d’extraction n’envoie pas les CV à un service externe.'),
    ('ACCÈS AUX DOCUMENTS', 'Extraction réservée à l’administrateur. CV, photos et certificats consultables par administrateur et superviseur.'),
    ('CONTRAINTES ACTUELLES', 'CV : 10 Mo ; pas de plafond fixe de pages dans le code actuel. L’OCR nécessite Windows et peut expirer sur un document long.')],
    'Toujours vérifier le texte et l’image avant de finaliser le dossier.',
    'Le README mentionne encore une ancienne limite de dix pages ; le code et le formulaire actuels n’imposent plus cette limite. Les PDF protégés doivent être déverrouillés. L’OCR dispose d’un budget de traitement d’environ 60 secondes. Les tests OCR utilisent une simulation, ils ne prouvent pas la disponibilité du moteur Windows de ce poste.', 'cv.py · CVImport.jsx · test_cv.py')

add('05 / PROJETS & ÉQUIPE', 'Préparer les projets et les affectations',
    'Administrateur et superviseur disposent des fonctions d’organisation.', [
    ('IDENTIFIER LE PROJET', 'Nom du projet, client, site industriel, description, bon de commande et nom du chef de projet client.'),
    ('AFFECTER LES TECHNICIENS', 'Rechercher et sélectionner un ou plusieurs techniciens actifs pour constituer l’équipe du projet.'),
    ('PARAMÉTRER LES TARIFS', 'Tarif client horaire, forfait client férié, coût technicien horaire et forfait technicien férié.'),
    ('PRÉPARER LA VALIDATION', 'Le nom du responsable ne crée pas un compte. L’administrateur configure séparément l’accès client du projet.')],
    'Un client actif et une affectation valide sont nécessaires pour les nouveaux pointages.',
    'Le formulaire actuel utilise un nom libre pour le chef de projet client. Il faut donc distinguer cette information descriptive d’un véritable compte de connexion. Le backend prend également en charge les comptes chefs de projet historiques rattachés au projet. Les champs financiers sont masqués dans les réponses projet des techniciens et clients.', 'ProjectSerializer · ProjectViewSet · forms.jsx · ProjectAccess.jsx')

add('05 / PROJETS & ÉQUIPE', 'L’espace du superviseur interne',
    'Coordonner le travail opérationnel et suivre les équipes.', [
    ('ORGANISER', 'Créer ou modifier les clients et projets, renseigner les informations commerciales et gérer les affectations.'),
    ('CONSULTER LES DOSSIERS', 'Accéder aux fiches techniciens, aux CV, aux photos, aux certificats et à la liste de l’équipe.'),
    ('SUIVRE LES INTERVENTIONS', 'Consulter les pointages, les projets actifs et les indicateurs opérationnels.'),
    ('CADRE DES AUTORISATIONS', 'La gestion des comptes, les accès client, la facturation, les encaissements et le calendrier restent réservés à l’administrateur.')],
    'Le superviseur prépare et suit ; le responsable client décide de la validation.',
    'Il peut voir les tarifs et les coûts calculés des projets. Il ne doit pas être présenté comme un simple lecteur de projets : le backend lui permet de créer et modifier les clients et projets. En revanche, il ne saisit pas les heures à la place du technicien.', 'views.py · access.py · App.jsx')

add('06 / ESPACE TECHNICIEN', 'Un accès centré sur le terrain',
    'Connexion habituelle ou page dédiée /tech/login.', [
    ('MES PROJETS', 'Consulter les projets affectés et retrouver les projets de son historique de pointages.'),
    ('MES HEURES', 'Choisir un projet actif, une date d’intervention, un nombre d’heures et une description des travaux.'),
    ('MON SUIVI', 'Consulter le détail de ses pointages, les motifs de refus et les actions disponibles pour corriger ou soumettre.'),
    ('MES GAINS', 'Consulter les heures validées ou facturées et les montants calculés pour soi, projet par projet.')],
    'Le technicien accède à ses propres pointages et à ses propres gains.',
    'Le formulaire actuel simplifie la saisie à une date et une durée. Il ne demande pas directement une heure de début, une heure de fin et une pause, même si ces informations existent dans le modèle serveur. Le technicien ne gère pas les comptes ni les documents privés des autres personnes.', 'App.jsx · forms.jsx · entries_for · projects_for')

add('06 / ESPACE TECHNICIEN', 'Le cycle d’un pointage',
    'Un état clair détermine les actions possibles.', kind='flow',
    steps=[('BROUILLON', 'Technicien', 'Saisir et corriger'),
           ('SOUMIS', 'Technicien', 'Transmettre au client'),
           ('VALIDÉ', 'Responsable client', 'Accepter les heures'),
           ('FACTURÉ', 'Administrateur', 'Intégrer à une facture')],
    cards=[('EN CAS DE REFUS', 'Le client renseigne un motif. Le technicien corrige le pointage refusé puis le soumet à nouveau.'),
           ('CONTRÔLES DE SAISIE', 'Projet actif et affectation vérifiés ; dates futures et chevauchements refusés ; durée contrôlée par le serveur.')],
    takeaway='Un pointage soumis, validé ou facturé est verrouillé à la modification.',
    notes='Les chevauchements sont recherchés sur tous les projets du technicien. Un refus crée un état distinct, modifiable par son auteur. Le client ne voit pas les brouillons. Les transitions et corrections sont historisées avec l’identité de l’acteur.', sources='TimeEntrySerializer · transition_entry · TimeEntryEvent')

add('06 / ESPACE TECHNICIEN', 'Les gains du technicien',
    'Le menu « Paiements » présente une estimation calculée à partir de l’activité validée.', [
    ('PÉRIMÈTRE PERSONNEL', 'Seules les heures validées ou facturées du technicien connecté sont prises en compte.'),
    ('CALCUL HORS JOURS FÉRIÉS', 'Les heures retenues sont multipliées par le tarif technicien configuré pour le projet.'),
    ('CALCUL DES JOURS FÉRIÉS', 'Un forfait technicien est appliqué une fois par technicien, date fériée et projet dans le calcul des gains.'),
    ('INFORMATION À INTERPRÉTER', 'Le montant affiché représente des gains calculés. Aucun statut de versement réel au technicien n’est enregistré par ce module.')],
    'Ces montants ne donnent pas accès aux factures ni aux encaissements du client.',
    'Les gains utilisent les tarifs et le calendrier actuels. Ils peuvent donc changer si ces paramètres sont modifiés. Ce module ne constitue pas un système de paie, un bulletin de salaire ou une preuve de virement. Les données de gains sont accessibles en lecture seule au technicien.', 'ProjectViewSet.earnings · ProjectSerializer.get_technician_labor_cost')

add('07 / ESPACE CLIENT', 'Deux modes d’accès au périmètre client',
    'Le rôle reste « Chef de projet client » ; le périmètre dépend du compte.', [
    ('COMPTE RESPONSABLE CLIENT', 'Connexion habituelle. Consultation des seuls projets attribués au compte, au sein de son entreprise cliente.'),
    ('COMPTE DÉDIÉ AU PROJET', 'Connexion via /project/login avec un identifiant créé par l’administrateur. Accès limité au projet associé.'),
    ('DONNÉES CONSULTABLES', 'Projets autorisés, pointages hors brouillons, informations limitées sur les techniciens du périmètre.'),
    ('DONNÉES NON EXPOSÉES', 'CV, photos et informations privées des techniciens ; tarifs internes, tarifs client, factures et paiements clients.')],
    'Un même client peut avoir plusieurs projets sans que tous ses comptes voient tout.',
    'Les restrictions reposent sur le projet attribué et le client actif. L’appartenance à la même entreprise ne suffit pas pour valider les pointages d’un autre projet. Le compte dédié est redirigé vers les pointages à valider.', 'projects_for · entries_for · LoginView · App.jsx')

add('07 / ESPACE CLIENT', 'Valider ou refuser les heures',
    'La décision client structure le passage vers la facturation.', [
    ('CONSULTER', 'Ouvrir les pointages soumis du projet autorisé et examiner le technicien, la date, la durée et les travaux.'),
    ('VALIDER', 'Accepter le pointage lorsqu’il correspond à l’intervention réalisée. Les heures deviennent éligibles à la facturation.'),
    ('REFUSER AVEC UN MOTIF', 'Préciser ce qui doit être corrigé. Le pointage retourne au technicien dans l’état Refusé.'),
    ('REPRENDRE LE SUIVI', 'Après correction et nouvelle soumission, le responsable client peut réexaminer le pointage.')],
    'Seuls les pointages à l’état Soumis peuvent être validés ou refusés.',
    'L’auteur de la décision, sa date et le motif éventuel sont enregistrés. Un administrateur ne dispose pas d’une action de validation de substitution dans le circuit actuel. La validation est une action métier tracée, pas une signature électronique certifiée.', 'transition_entry · entry-detail dans App.jsx')

add('07 / ESPACE CLIENT', 'Configurer un accès client par projet',
    'Parcours administrateur : Projets → Accès client.', [
    ('CRÉER L’ACCÈS', 'Choisir l’identifiant du projet, saisir ou générer un mot de passe et activer le compte.'),
    ('TRANSMETTRE LES INFORMATIONS', 'Communiquer l’adresse de connexion /project/login et les identifiants au contact autorisé.'),
    ('MAINTENIR L’ACCÈS', 'Modifier l’identifiant, remplacer le mot de passe ou désactiver le compte si nécessaire.'),
    ('LIMITER LE PÉRIMÈTRE', 'Le compte est associé à un projet unique. Il ne donne pas automatiquement accès aux autres projets du client.')],
    'Le mot de passe stocké ne peut pas être relu : le noter lors de sa création.',
    'Le formulaire permet de générer un mot de passe. L’envoi d’un email ou d’une invitation automatique n’est pas implémenté dans ce parcours. Un client désactivé ne peut plus utiliser cet accès projet.', 'ProjectAccess.jsx · ProjectViewSet.client_access · LoginView')

add('08 / SUIVI FINANCIER', 'Facturer les interventions validées',
    'Fonction réservée à l’administrateur.', [
    ('SÉLECTIONNER', 'Choisir un projet et ses pointages validés, encore non facturés.'),
    ('PARAMÉTRER', 'Vérifier l’échéance, les tarifs applicables et le taux de taxe saisi explicitement.'),
    ('ÉMETTRE', 'Créer une facture numérotée avec ses lignes, montants et informations de référence figés à l’émission.'),
    ('CONSULTER ET IMPRIMER', 'Ouvrir le détail de la facture, consulter ses paiements et imprimer depuis le navigateur.')],
    'Un pointage facturé ne peut pas être sélectionné une seconde fois.',
    'Le serveur bloque les pointages non validés, les doublons et le mélange de projets. La facture conserve notamment les noms client et projet, l’adresse, le bon de commande et les montants. Le document imprimable actuel est décrit comme un outil de suivi interne ; les informations de l’entreprise émettrice restent à compléter avant usage commercial définitif.', 'create_invoice · InvoiceCreateSerializer · InvoiceDetail · README')

add('08 / SUIVI FINANCIER', 'Suivre les règlements clients',
    'Un enregistrement distinct pour chaque paiement.', [
    ('ENREGISTRER', 'Sélectionner la facture, renseigner le montant, la date, le mode de paiement et la référence.'),
    ('MODES DISPONIBLES', 'Virement, espèces ou chèque. Le paiement peut être partiel ou couvrir le solde restant.'),
    ('SUIVRE LE SOLDE', 'La facture évolue entre Non payée, Partiellement payée et Payée selon les montants enregistrés.'),
    ('CONTRÔLER LES MONTANTS', 'Montant positif, pas de dépassement du solde, pas de date future ni de date antérieure à la facture.')],
    'Un paiement enregistré met à jour le solde utilisé dans le suivi administrateur.',
    'Ces enregistrements concernent les paiements reçus des clients. Aucun raccordement bancaire automatique n’est présent dans le périmètre examiné. Les modifications et suppressions de paiements ne sont pas exposées par l’API actuelle.', 'record_payment · PaymentSerializer · Invoice.balance')

add('08 / SUIVI FINANCIER', 'Tarifs et jours fériés',
    'Distinguer le montant facturé au client du coût calculé pour le technicien.', [
    ('JOUR ORDINAIRE — CLIENT', 'Heures validées sélectionnées × tarif client horaire du projet.'),
    ('JOUR ORDINAIRE — TECHNICIEN', 'Heures personnelles validées ou facturées × tarif technicien horaire du projet.'),
    ('JOUR FÉRIÉ — CLIENT', 'Forfait client par technicien et date fériée, réparti entre les pointages sélectionnés dans la facture.'),
    ('JOUR FÉRIÉ — TECHNICIEN', 'Forfait technicien par date fériée travaillée et projet ; calendrier géré par l’administrateur.')],
    'Les forfaits client et technicien sont deux paramètres indépendants.',
    'Le calcul du forfait client regroupe les pointages sélectionnés dans une seule émission de facture. Il ne faut pas promettre une déduplication du forfait entre plusieurs factures séparées portant sur la même date et le même technicien. Les dates fériées sont saisies par l’administrateur, sans présumer d’un calendrier légal automatique.', 'create_invoice · get_technician_labor_cost · PublicHolidayViewSet')

add('09 / CONFIANCE & ACCÈS', 'Séparer les données et conserver les traces',
    'Des contrôles intégrés au fonctionnement métier.', [
    ('AUTHENTIFICATION', 'Sessions Django, mots de passe hachés, protection CSRF et limitation des tentatives de connexion.'),
    ('ISOLATION DES PÉRIMÈTRES', 'Filtrage des projets et pointages selon le rôle, l’affectation et le compte client associé.'),
    ('DOCUMENTS PRIVÉS', 'Téléchargement authentifié des CV, photos et certificats ; accès limité aux rôles internes autorisés.'),
    ('TRAÇABILITÉ', 'Historique daté des états de pointage, auteur des décisions, motif de refus et référence de facture.')],
    'Ces mécanismes sont implémentés ; ils ne constituent pas une certification de sécurité.',
    'La désactivation permet de conserver les liens historiques. Les contrôles du backend complètent l’affichage conditionnel des menus. Cette présentation ne remplace pas un audit de sécurité ou une recette complète d’exploitation.', 'access.py · views.py · services.py · serializers.py')

add('09 / CONFIANCE & ACCÈS', 'Un usage web et un accès Android',
    'L’interface s’adapte aux écrans de bureau et de téléphone.', [
    ('NAVIGATEUR', 'Accès à l’application web en français depuis un ordinateur ou un téléphone.'),
    ('ESPACE TECHNICIEN', 'Présentation adaptée au mobile pour les projets, les pointages et les gains personnels.'),
    ('ANDROID LOCAL', 'Une application de développement ouvre l’espace technicien dans une WebView avec identité Copetrolite.'),
    ('CONDITIONS ACTUELLES', 'Le paquet Android documenté dépend du réseau local et des serveurs démarrés. Un accès public demande une configuration de déploiement.')],
    'La version Android actuelle n’est pas présentée comme une application hors ligne.',
    'Le README décrit un APK de développement lié à une adresse réseau locale. Une distribution plus large nécessite notamment une adresse accessible en HTTPS et une signature de production. Ces travaux ne sont pas réalisés dans le cadre de cette présentation.', 'README · frontend/src/App.jsx · styles.css')

add('10 / SOCLE TECHNIQUE', 'Une architecture organisée par responsabilité',
    'Vue de la configuration réellement présente dans le projet.', kind='flow',
    steps=[('INTERFACE', 'React / Vite', 'Écrans et formulaires'),
           ('API MÉTIER', 'Django REST', 'Accès et validations'),
           ('DONNÉES', 'MySQL configuré', 'Données relationnelles'),
           ('DOCUMENTS', 'Stockage local', 'CV, photos, certificats')],
    cards=[('EXTRACTION DOCUMENTAIRE', 'PyMuPDF lit les PDF et les images. L’OCR Windows fournit le texte des pages scannées.'),
           ('COHÉRENCE MÉTIER', 'Les règles d’accès, les transitions de pointage et les calculs financiers sont exécutés côté serveur.')],
    takeaway='Les tests rapides utilisent une base SQLite en mémoire, distincte de la base applicative.',
    notes='Les consignes générales mentionnent PostgreSQL, mais coproject/settings.py configure actuellement django.db.backends.mysql et le README documente MySQL. La présentation reflète cette configuration observée, sans la modifier. La base de données et les documents locaux doivent être sauvegardés ensemble.', sources='coproject/settings.py · test_settings.py · cv.py · README')

add('11 / DÉMONSTRATION', 'Un exemple simple, de l’heure au solde',
    'Exemple fictif, hors jour férié, avec un taux de taxe de 0 % pour simplifier.', kind='example',
    notes='Exemple purement pédagogique : huit heures sont déclarées puis validées. Avec un tarif client de 2 500 DZD par heure, la facture hors taxe atteint 20 000 DZD. Avec un tarif technicien de 1 000 DZD, ses gains calculés sont de 8 000 DZD. Après un règlement client de 12 000 DZD, le solde restant est de 8 000 DZD. Aucune donnée réelle du projet n’est utilisée.',
    sources='Formules : services.py · ProjectSerializer')

add('11 / DÉMONSTRATION', 'Le scénario à présenter en direct',
    'Un parcours de 8 à 10 minutes avec des données de démonstration.', [
    ('1 · ADMINISTRATEUR', 'Présenter un client et un dossier technicien ; extraire l’expérience et la photo d’un CV, relire et enregistrer.'),
    ('2 · ORGANISATION', 'Créer un projet, affecter le technicien et configurer un accès client dédié.'),
    ('3 · TECHNICIEN ET CLIENT', 'Saisir puis soumettre les heures ; se connecter côté client pour valider ou refuser avec un motif.'),
    ('4 · ADMINISTRATEUR', 'Créer la facture à partir du pointage validé, enregistrer un paiement et montrer le solde restant.')],
    'Préparer un CV de démonstration et des comptes de test avant la présentation.',
    'Changer explicitement de session entre les rôles pour montrer les différences de menus et de données. Ne pas manipuler de véritables données personnelles dans une démonstration publique. Préparer un CV lisible avec un portrait intégré pour rendre le résultat de l’extraction observable.', 'Parcours issus des écrans et règles métier existants')

add('12 / ÉTAT DU PROJET', 'Périmètre livré et vérifications',
    'État observé le 18 octobre 2026 sur le code local.', [
    ('PARCOURS MÉTIER', '35 tests métier et accès réussis : rôles, isolation, validation, factures, paiements et gains.'),
    ('CV ET PROFILS', '26 tests exécutés : 23 réussis et 3 en échec. Les écarts concernent un format d’adresse et des scénarios de mise à jour de profil.'),
    ('FONCTIONS PRÉSENTÉES', 'Gestion des équipes et projets, extraction CV et photo, pointages, validation client et suivi financier présents dans le code.'),
    ('AVANT DIFFUSION LARGE', 'Résoudre les écarts de tests, compléter les informations de l’émetteur sur les factures et préparer le déploiement public.')],
    'Les essais utilisent SQLite en mémoire ; ils ne constituent pas une recette de production.',
    'Tests lancés séparément : eliteapp.tests puis eliteapp.test_cv avec coproject.test_settings. Total : 61 tests, 58 réussis, 2 échecs d’assertion et 1 erreur. Les écarts CV/profils sont préexistants : attente de casse de l’adresse, validation du téléphone lors de la mise à jour photo, absence de certificate_documents dans une réponse de mise à jour. Le moteur OCR est simulé dans ses tests. Aucun code applicatif n’a été modifié pour cette livraison documentaire.', 'Exécution locale : manage.py test · 18/10/2026')

add('12 / CONCLUSION', 'Une chaîne de travail partagée',
    'De la compétence du technicien au règlement de l’intervention.', kind='closing',
    cards=[('ADMINISTRATEUR', 'Structure les accès et suit les finances.'),
           ('SUPERVISEUR', 'Organise les projets et les équipes.'),
           ('TECHNICIEN', 'Renseigne les interventions et suit ses gains.'),
           ('CLIENT', 'Vérifie les heures et décide de leur validation.')],
    takeaway='COPETROLITE · Une responsabilité claire à chaque étape.',
    notes='Conclure sur la complémentarité des rôles. L’extraction de l’expérience et de la photo facilite la constitution des dossiers techniciens ; le circuit de pointage structure ensuite le suivi du travail. Ouvrir les échanges sur les besoins de déploiement et les priorités d’évolution.')

add('ANNEXE / RÉFÉRENCES', 'Sources et limites de cette présentation',
    'Présentation établie à partir du projet local ; aucune donnée métier réelle n’est reproduite.', [
    ('RÔLES ET CIRCUIT MÉTIER', 'eliteapp/models.py, access.py, views.py, serializers.py et services.py.'),
    ('ÉCRANS ET DOCUMENTS', 'App.jsx, forms.jsx, ProjectAccess.jsx, CVImport.jsx, ExperienceEditor.jsx et TechnicianPhoto.jsx.'),
    ('EXTRACTION ET VÉRIFICATION', 'eliteapp/cv.py, scripts/ocr-image.ps1, eliteapp/tests.py et eliteapp/test_cv.py.'),
    ('CONFIGURATION ET IDENTITÉ', 'coproject/settings.py, test_settings.py, README.md ; logo existant frontend/src/assets/CO-1-1-1.png.')],
    'Les comportements du code actuel priment sur les descriptions anciennes du README.',
    'Écarts documentaires identifiés : ancienne limite de dix pages du CV, anciennes informations sur le matricule, et ancienne description de la saisie horaire. Les fonctionnalités sont présentées selon les fichiers actuels. L’environnement utilise MySQL malgré la mention générique de PostgreSQL dans les consignes. Les fichiers de configuration sensibles et les données applicatives n’ont pas été lus pour produire ce document.', 'Sources locales uniquement')


# The requested short edition keeps exactly fifteen slides, with actual UI captures.
long_version = {s['title']: s for s in SLIDES}
SLIDES = []
def keep(title):
    SLIDES.append(long_version[title])

def screen(section, title, subtitle, pictures, points, takeaway, layout='hero'):
    add(section, title, subtitle, kind='screen', pictures=pictures, points=points,
        takeaway=takeaway, layout=layout,
        notes='Capture du véritable frontend React connecté au backend Django, dans une base SQLite de démonstration séparée. Les comptes, entreprises, montants et documents sont fictifs ; le portrait est un avatar illustratif. ' + ' '.join(t+' : '+b for t,b in points),
        sources='Captures réelles de l’application · Données de démonstration · 18/10/2026')

keep('COPETROLITE')
keep('Le parcours métier de bout en bout')
keep('Matrice des droits')
screen('01 / ADMINISTRATEUR', 'Piloter l’activité en un regard',
    'Tableau de bord administrateur : projets, heures et encaissements.',
    [('02-tableau-bord-admin.png', (0,0,1,.72))],
    [('ACTIVITÉ', 'Retrouver les projets actifs et les heures pointées.'),
     ('VALIDATION', 'Identifier les pointages en attente de décision client.'),
     ('FINANCES', 'Suivre le montant restant à encaisser sur les factures.')],
    'L’administrateur dispose de la vue globale et gère les comptes et les finances.')
screen('02 / ÉQUIPE & PROJETS', 'Organiser les personnes et les missions',
    'En haut : Équipe côté administrateur. En bas : Projets côté superviseur.',
    [('04-equipe.png',None),('11-projets-superviseur.png',None)],
    [('ADMINISTRATEUR', 'Créer les comptes, choisir les rôles et activer les accès.'),
     ('SUPERVISEUR', 'Gérer les clients, les projets et les affectations.'),
     ('PÉRIMÈTRE', 'Le superviseur consulte les dossiers ; il ne gère pas les comptes.')],
    'Chaque projet relie un client, un site, un responsable et une équipe.', 'dualwide')
screen('03 / DOSSIER TECHNICIEN', 'Une fiche structurée pour chaque technicien',
    'Identité et affectation dans le profil ; CV, expérience et certificats plus bas dans la fiche.',
    [('05-profil-technicien.png',(0,0,1,.55))],
    [('PROFIL', 'Photo, spécialité, coordonnées et matricule automatique.'),
     ('COMPÉTENCES', 'Expérience professionnelle, CV et certificats joints.'),
     ('ACCÈS', 'Modification par l’administrateur ; consultation par le superviseur.')],
    'Les documents privés du technicien ne sont pas exposés au client.')
screen('04 / EXTRACTION DU CV', 'Extraire l’expérience professionnelle',
    'Résultat réel obtenu depuis le CV de démonstration.',
    [('06-extraction-experience.png',None)],
    [('IMPORTER', 'CV PDF, JPG ou PNG, jusqu’à 10 Mo.'),
     ('EXTRAIRE', 'Lecture du PDF ; OCR Windows pour les images et scans.'),
     ('VÉRIFIER', 'Relire, corriger, utiliser le texte puis enregistrer la fiche.')],
    'L’extraction prépare le contenu ; le bouton Enregistrer le sauvegarde.')
screen('04 / EXTRACTION DU CV', 'Récupérer la photo du CV',
    'Aperçu après l’action « Utiliser la photo du CV » — portrait illustratif.',
    [('07-photo-extraite.png',None)],
    [('DÉTECTER', 'Le système propose une image candidate du document.'),
     ('CONTRÔLER', 'Vérifier l’aperçu ou choisir une photo séparée.'),
     ('ENREGISTRER', 'La photo sélectionnée est sauvegardée avec la fiche.')],
    'La sélection de l’image est heuristique : la vérification visuelle reste nécessaire.')
screen('05 / ESPACE TECHNICIEN', 'Déclarer ses interventions sur mobile',
    'Deux écrans réels : l’accueil et le formulaire de pointage.',
    [('13-technicien-mobile.png',None),('14-pointage-mobile.png',None)],
    [('MES PROJETS', 'Consulter les missions affectées et les derniers pointages.'),
     ('MES HEURES', 'Choisir le projet, la date, la durée et décrire les travaux.'),
     ('MON SUIVI', 'Soumettre les heures et corriger un pointage refusé.')],
    'Le technicien voit ses propres pointages ; les heures soumises sont verrouillées.', 'mobilepair')
screen('06 / ESPACE CLIENT', 'Valider les heures du projet',
    'Le client voit les pointages soumis de son périmètre.',
    [('12-validation-client.png',None)],
    [('CONSULTER', 'Vérifier le technicien, la date et les heures.'),
     ('DÉCIDER', 'Valider ou refuser avec un motif obligatoire.'),
     ('TRACER', 'Conserver l’auteur, la date et l’historique.')],
    'Un refus permet la correction et une nouvelle soumission par le technicien.', 'wide')
screen('06 / ESPACE CLIENT', 'Un accès client limité à chaque projet',
    'Configuration administrateur depuis la liste des projets.',
    [('08-acces-client.png',None)],
    [('CRÉATION', 'Choisir un identifiant et un mot de passe pour le projet.'),
     ('CONNEXION', 'Le client utilise la page dédiée /project/login.'),
     ('CONTRÔLE', 'Remplacer le mot de passe ou désactiver l’accès.')],
    'Le compte dédié ne donne pas accès aux autres projets de l’entreprise cliente.')
screen('07 / FACTURATION', 'Transformer les heures validées en facture',
    'Exemple réel de facture dans l’environnement de démonstration.',
    [('09-facture.png',None)],
    [('ÉMETTRE', 'Sélectionner les pointages validés et non encore facturés.'),
     ('CONSERVER', 'Les informations et montants sont figés à l’émission.'),
     ('SUIVRE', 'Exemple : 40 000 DZD facturés, 25 000 reçus, solde 15 000.')],
    'Facturation réservée à l’administrateur ; consultation et impression disponibles.')
screen('07 / PAIEMENTS & GAINS', 'Deux suivis financiers distincts',
    'À gauche : encaissement client. À droite : gains calculés du technicien.',
    [('10-paiements-client.png',None),('15-gains-technicien.png',None)],
    [('ADMINISTRATEUR', 'Enregistrer un paiement partiel ou complet et contrôler le solde.'),
     ('TECHNICIEN', 'Consulter ses gains par projet, sur ses heures validées ou facturées.')],
    'Les gains calculés ne constituent pas une preuve de versement au technicien.', 'finance')
add('08 / FIABILITÉ', 'Des contrôles clairs, un périmètre connu',
    'Les fonctionnalités sont présentes dans le code ; les points de validation restent visibles.', [
    ('PROTECTION DES ACCÈS', 'Sessions, mots de passe hachés, protection CSRF et filtrage serveur selon le rôle et le projet.'),
    ('TRAÇABILITÉ', 'Historique des pointages, refus motivés et prévention de la double facturation d’un même pointage.'),
    ('VÉRIFICATIONS', '35 tests métier réussis ; 23 tests CV/profils réussis sur 26. Trois écarts préexistants restent à résoudre.'),
    ('MISE EN SERVICE', 'Compléter les informations émetteur des factures et préparer le déploiement HTTPS et les sauvegardes.')],
    'Captures sur données fictives. Le CV et la photo sont traités localement.',
    notes=long_version['Périmètre livré et vérifications']['notes'],
    sources='Règles et tests du projet local · Captures issues d’une base de démonstration isolée')
keep('Une chaîne de travail partagée')
assert len(SLIDES) == 15

pdfmetrics.registerFont(TTFont('ArialLocal', 'C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('ArialLocalBold', 'C:/Windows/Fonts/arialbd.ttf'))
prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
prs.core_properties.title = 'Copetrolite — Présentation du projet'
prs.core_properties.subject = 'Accès, gestion des interventions, extraction CV et photo'
prs.core_properties.author = 'Copetrolite'
prs.core_properties.keywords = 'Copetrolite, interventions, techniciens, CV, expérience, photo'
pdf = canvas.Canvas(str(OUT / 'Copetrolite_Presentation_FR.pdf'), pagesize=(W, H))
pdf.setTitle('Copetrolite — Présentation du projet')
pdf.setAuthor('Copetrolite')
pdf.setSubject('Présentation fonctionnelle détaillée en français')
slide = None
layout_log = []

def rect(x, y, w, h, color, radius=0):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                Pt(x), Pt(y), Pt(w), Pt(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    if radius:
        sh.adjustments[0] = min(radius / min(w, h), 0.5)
    pdf.setFillColor(HexColor('#' + color))
    if radius:
        pdf.roundRect(x, H-y-h, w, h, radius, fill=1, stroke=0)
    else:
        pdf.rect(x, H-y-h, w, h, fill=1, stroke=0)
    return sh

def wrap(value, width, size, bold=False):
    font = 'ArialLocalBold' if bold else 'ArialLocal'
    lines = []
    for para in value.split('\n'):
        line = ''
        for word in para.split():
            proposal = f'{line} {word}'.strip()
            if pdfmetrics.stringWidth(proposal, font, size) <= width:
                line = proposal
            else:
                if not line:
                    raise ValueError(f'Word too wide: {word}')
                lines.append(line)
                line = word
        lines.append(line)
    return lines

def text(value, x, y, w, size=18, color=INK, bold=False, max_h=None, leading=None):
    lines = wrap(value, w, size, bold)
    leading = leading or size * 1.24
    height = len(lines) * leading
    if max_h is not None and height > max_h + 0.1:
        raise ValueError(f'Slide {len(prs.slides)} overflows ({height:.1f}>{max_h}): {value}')
    if y + height > H or x + w > W + 1 or x < 0 or y < 0:
        raise ValueError(f'Outside page: {value}')
    font = 'ArialLocalBold' if bold else 'ArialLocal'
    for i, line in enumerate(lines):
        top = y + i * leading
        tx = slide.shapes.add_textbox(Pt(x), Pt(top), Pt(w + 3), Pt(size * 1.5))
        tf = tx.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.space_before = p.space_after = Pt(0)
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.text = line
        pdf.setFillColor(HexColor('#' + color))
        pdf.setFont(font, size)
        pdf.drawString(x, H-top-size*0.96, line)
    layout_log.append(dict(page=len(prs.slides), text=value, x=x, y=y, width=w, height=height))
    return height

def logo(x, y, w):
    h = w * 296 / 419
    slide.shapes.add_picture(str(LOGO), Pt(x), Pt(y), width=Pt(w), height=Pt(h))
    pdf.drawImage(str(LOGO), x, H-y-h, width=w, height=h, mask='auto')

def card(title, body, x, y, w=426, h=128, dark=False):
    rect(x, y, w, h, '12525A' if dark else WHITE, 8)
    rect(x, y+16, 3, h-32, GOLD if dark else GREEN)
    text(title, x+20, y+15, w-38, 14, MINT if dark else GREEN, True, max_h=36)
    text(body, x+20, y+43, w-40, 16, WHITE if dark else INK, max_h=h-51)

def footer(s, n, dark=False):
    color = 'B2C9CA' if dark else MUTED
    if s.get('sources'):
        text(s['sources'], 42, 512, 780, 8, color, max_h=13)
    text(f'COPETROLITE  /  {n:02d}', 823, 512, 110, 8, color, True)

def normal_header(s):
    rect(0, 0, W, H, BG)
    rect(0, 0, W, 63, PETROL)
    logo(31, -0.0, 85)
    text(s['section'], 140, 25, 710, 11, MINT, True)
    text(s['title'], 42, 84, 880, 30, PETROL, True, max_h=75)
    text(s['subtitle'], 44, 132, 876, 13.5, MUTED, max_h=35)

def screenshot(spec, x, y, w, h):
    filename, crop = spec
    path = OUT / 'captures' / filename
    iw, ih = Image.open(path).size
    left, top, right, bottom = crop or (0,0,1,1)
    aspect = iw*(right-left)/(ih*(bottom-top))
    dw, dh = (w,w/aspect) if w/h < aspect else (h*aspect,h)
    dx, dy = x+(w-dw)/2, y+(h-dh)/2
    rect(dx-2,dy-2,dw+4,dh+4,LINE,3)
    pic = slide.shapes.add_picture(str(path),Pt(dx),Pt(dy),width=Pt(dw),height=Pt(dh))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = left, top, 1-right, 1-bottom
    pdf.saveState()
    clip = pdf.beginPath();clip.rect(dx,H-dy-dh,dw,dh);pdf.clipPath(clip,stroke=0)
    fw, fh = dw/(right-left),dh/(bottom-top)
    pdf.drawImage(str(path),dx-left*fw,H-dy-fh+top*fh,width=fw,height=fh,mask='auto')
    pdf.restoreState()

def screen_points(points,x=670,y=190,w=244,step=87):
    for i,(title,body) in enumerate(points):
        text(title,x,y+i*step,w,12,GREEN,True,max_h=30)
        text(body,x,y+23+i*step,w,15,INK,max_h=58)

for n, s in enumerate(SLIDES, 1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pdf.bookmarkPage(f'page-{n}')
    pdf.addOutlineEntry(f'{n:02d} · {s["title"]}', f'page-{n}', level=0)
    slide.notes_slide.notes_text_frame.text = s['notes'] + ('\n\nSources : ' + s['sources'] if s['sources'] else '')
    kind = s['kind']
    if kind == 'cover':
        rect(0, 0, W, H, PETROL)
        rect(666, 0, 294, H, '0D5057')
        rect(42, 171, 64, 5, GOLD)
        logo(40, 15, 174)
        text('PRÉSENTATION DU PROJET', 44, 148, 560, 12, MINT, True)
        text('COPETROLITE', 40, 194, 630, 55, WHITE, True)
        text(s['subtitle'], 44, 270, 570, 28, WHITE, max_h=75)
        text(s['takeaway'], 44, 370, 540, 19, MINT, max_h=55)
        text('ADMINISTRATION · ÉQUIPE\nTECHNICIENS · CLIENTS', 44, 453, 600, 12, WHITE, True)
        text('SYNTHÈSE · 15 DIAPOSITIVES · CAPTURES DE L’APPLICATION', 44, 505, 620, 9, 'B2C9CA')
        for i, (num, label) in enumerate([('01', 'Structurer'), ('02', 'Intervenir'), ('03', 'Valider'), ('04', 'Facturer')]):
            yy = 116 + i * 85
            text(num, 703, yy, 80, 30, GOLD, True)
            text(label, 767, yy+8, 170, 20, WHITE)
    elif kind == 'closing':
        rect(0, 0, W, H, PETROL)
        logo(33, 3, 112)
        text(s['title'], 42, 105, 880, 37, WHITE, True)
        text(s['subtitle'], 44, 158, 870, 18, MINT)
        for i, (title, body) in enumerate(s['cards']):
            card(title, body, 42+(i%2)*450, 218+(i//2)*111, 426, 95, True)
        text(s['takeaway'], 44, 465, 870, 20, WHITE, True)
        footer(s, n, True)
    else:
        normal_header(s)
        if kind == 'screen':
            layout=s['layout']
            if layout=='hero':
                screenshot(s['pictures'][0],42,177,598,278)
                screen_points(s['points'])
            elif layout=='dualwide':
                screenshot(s['pictures'][0],42,177,598,133)
                screenshot(s['pictures'][1],42,324,598,133)
                screen_points(s['points'])
            elif layout=='mobilepair':
                screenshot(s['pictures'][0],70,174,190,282)
                screenshot(s['pictures'][1],290,174,190,282)
                screen_points(s['points'],x=535,w=375)
            elif layout=='wide':
                screenshot(s['pictures'][0],42,176,876,174)
                for i,(title,body) in enumerate(s['points']):
                    card(title,body,42+i*298,361,280,96)
            elif layout=='finance':
                screenshot(s['pictures'][0],42,176,597,135)
                screen_points([s['points'][0]],x=52,y=335,w=570)
                screenshot(s['pictures'][1],673,171,220,285)
                text('Gains personnels',665,451,245,11,GREEN,True)
        elif kind == 'cards':
            for i, (title, body) in enumerate(s['cards']):
                card(title, body, 42+(i%2)*450, 183+(i//2)*141)
        elif kind == 'flow':
            steps = s['steps']
            gap = 16
            cw = (876 - (len(steps)-1)*gap)/len(steps)
            for i, (title, who, body) in enumerate(steps):
                x = 42+i*(cw+gap)
                rect(x, 187, cw, 124, PETROL, 7)
                text(f'{i+1:02d}', x+14, 198, cw-24, 19, GOLD, True)
                text(title, x+14, 229, cw-24, 11.5, WHITE, True, max_h=29)
                text(who, x+14, 253, cw-24, 10.5, MINT, max_h=26)
                text(body, x+14, 279, cw-24, 9.5, WHITE, max_h=28)
                if i < len(steps)-1:
                    text('›', x+cw+3, 228, 12, 22, GREEN, True)
            for i, (title, body) in enumerate(s['cards']):
                card(title, body, 42+i*450, 331, 426, 122)
        elif kind == 'matrix':
            widths = [263, 133, 149, 179, 152]
            y = 181
            for rowidx, row in enumerate([s['headers']] + s['rows']):
                rh = 25.5
                x = 42
                for col, value in enumerate(row):
                    rect(x, y, widths[col]-1, rh-1, PETROL if rowidx == 0 else (WHITE if rowidx%2 else 'E5EEEB'))
                    text(value, x+9, y+6, widths[col]-16, 11.3 if rowidx else 12,
                         WHITE if rowidx == 0 else INK, rowidx == 0 or col == 0, max_h=20)
                    x += widths[col]
                y += rh
        elif kind == 'example':
            metrics = [('8 h', 'Pointage validé'), ('20 000 DZD', 'Facture client'), ('8 000 DZD', 'Gains calculés')]
            for i, (value, label) in enumerate(metrics):
                x = 42 + i * 300
                rect(x, 187, 276, 113, PETROL, 8)
                text(value, x+18, 207, 242, 30, WHITE, True)
                text(label, x+18, 258, 242, 15, MINT)
            card('PARAMÈTRES FICTIFS', 'Tarif client : 2 500 DZD / h\nTarif technicien : 1 000 DZD / h\nTaxe de démonstration : 0 %', 42, 321, 426, 129)
            card('RÈGLEMENT CLIENT PARTIEL', 'Facture : 20 000 DZD\nPaiement reçu : 12 000 DZD\nReste à encaisser : 8 000 DZD', 492, 321, 426, 129)
            s['takeaway'] = 'Les gains calculés du technicien et le solde client sont deux informations distinctes.'
        if s['takeaway']:
            rect(42, 470, 876, 32, 'E0EEE8', 5)
            text(s['takeaway'], 55, 478, 850, 12, PETROL, True, max_h=23)
        footer(s, n)
    pdf.showPage()

pdf.save()
prs.save(OUT / 'Copetrolite_Presentation_FR.pptx')
(OUT / 'contenu_presentation.json').write_text(json.dumps(SLIDES, ensure_ascii=False, indent=2), encoding='utf-8')
(OUT / 'controle_mise_en_page.json').write_text(json.dumps(layout_log, ensure_ascii=False, indent=2), encoding='utf-8')
guide = ['# Copetrolite — Guide de présentation', '',
         f'Support en français : {len(SLIDES)} diapositives au format 16:9, avec logo et captures réelles du projet.',
         'Durée suggérée : 12 à 15 minutes. Les captures utilisent des données fictives dans une base de démonstration séparée.', '',
         'Le PowerPoint est modifiable. Les notes ci-dessous figurent aussi dans ses notes de présentation.', '']
for i, s in enumerate(SLIDES, 1):
    guide += [f'## {i:02d}. {s["title"]}', '', s['notes'], '', 'Sources : ' + (s['sources'] or 'Synthèse du projet'), '']
guide += ['## Vérifications exécutées', '',
          '- `env/Scripts/python.exe manage.py test eliteapp.tests --settings=coproject.test_settings --verbosity 1` : 35 tests réussis.',
          '- `env/Scripts/python.exe manage.py test eliteapp.test_cv --settings=coproject.test_settings --verbosity 1` : 26 tests, 23 réussis, 2 échecs et 1 erreur.',
          '- Écarts préexistants : attente de casse de l’adresse ; téléphone non conforme lors du scénario de photo ; réponse sans certificate_documents lors du scénario de mise à jour des certificats.',
          '- Le moteur OCR est simulé dans les tests ; aucune recette OCR réelle ou de production n’est revendiquée.',
          '- Aucune modification du code applicatif, de la configuration ou de la base métier.', '',
          '## Régénération', '',
          'Exécuter `python presentation/generate_presentation.py` depuis la racine. Le générateur utilise ReportLab disponible sur le poste et python-pptx installé séparément dans `.runtime/presentation-tools`. Ces outils ne sont pas ajoutés aux dépendances de l’application.', '']
(OUT / 'Guide_de_presentation_FR.md').write_text('\n'.join(guide), encoding='utf-8')
print(f'Generated {len(SLIDES)} slides; all text boxes passed size checks.')
